"""Compile one retained composition, preserving representable source content."""
from functools import lru_cache
from io import BytesIO
import os
from pathlib import Path
import re

import fitz
from PIL import Image, ImageOps
from pypdf import PdfReader, PdfWriter, Transformation
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from ..transport.errors import StudioError
from .artifacts import document_bytes
from .boards import BoardExport
from .publications import read_publication


@lru_cache(maxsize=1)
def _font():
    # Use one installed face for measurement, embedded PDF and PPTX. Prefer a
    # CJK face on every platform so the result does not depend on whether the
    # first text encountered happens to be Latin. Linux CI/runtime images
    # install fonts-wqy-microhei; Windows packages use Microsoft YaHei.
    from .drawings import _sheet_fonts
    candidates = (
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/msyh.ttc",
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
    )
    path = next((candidate for candidate in candidates if candidate.is_file()), None)
    if path is None:
        path = _sheet_fonts()["normal"]
    font = TTFont("AF_Publication", str(path), subfontIndex=0)
    pdfmetrics.registerFont(font)
    return font, font.face.familyName.decode("utf-8")


def _lines(text, width, size):
    font, _ = _font()
    if any(ord(character) not in font.face.charToGlyph for character in text if character != "\n"):
        raise StudioError(422, "PUBLICATION_FONT_GLYPH", "The publication font cannot display this text. Use supported characters or install a compatible font.")
    lines = []
    for paragraph in text.split("\n"):
        current = ""
        for token in re.findall(r"[A-Za-z0-9_]+|[^A-Za-z0-9_]", paragraph):
            if current and pdfmetrics.stringWidth(current + token, font.fontName, size) > width:
                lines.append(current.rstrip())
                current = ""
            for character in token:
                if pdfmetrics.stringWidth(character, font.fontName, size) > width:
                    raise StudioError(422, "PUBLICATION_TEXT_OVERFLOW", "A character is wider than its text box. Enlarge the box or reduce the font size.")
                if current and pdfmetrics.stringWidth(current + character, font.fontName, size) > width:
                    lines.append(current)
                    current = ""
                current += character
        lines.append(current.rstrip())
    return lines


def _placement(item, width, height):
    scale = min(item["width"] / width, item["height"] / height)
    w, h = width * scale, height * scale
    return item["x"] + (item["width"] - w) / 2, item["y"] + (item["height"] - h) / 2, w, h, scale


def _image(data, mime_type, item):
    if mime_type == "application/pdf":
        with fitz.open(stream=data, filetype="pdf") as document:
            page = document[item["source"]["pageIndex"]]
            scale = min(2, 2048 / max(page.rect.width, page.rect.height))
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=True)
            # MuPDF pixels use premultiplied alpha; PNG encoding unpremultiplies
            # before Pillow and PowerPoint consume the image.
            data = pixmap.tobytes("png")
    with Image.open(BytesIO(data)) as source:
        original = ImageOps.exif_transpose(source).convert("RGBA")
        original.thumbnail((2048, 2048), Image.Resampling.LANCZOS)
        left, top, right, bottom = item["crop"]
        box = (round(left * original.width), round(top * original.height), round((1 - right) * original.width), round((1 - bottom) * original.height))
        if box[0] >= box[2] or box[1] >= box[3]:
            raise StudioError(422, "PUBLICATION_CROP", "The crop leaves no visible pixels.")
        image = original.crop(box)
        placement = _placement(item, image.width, image.height)
        output = BytesIO()
        image.save(output, format="PNG")
    return output.getvalue(), placement[:4]


def _pdf_source(target, data, item, page_height):
    reader = PdfReader(BytesIO(data))
    source = reader.pages[item["source"]["pageIndex"]]
    # Page annotations are outside the content-stream crop, and document layer
    # states / page isolation groups cannot be transferred by merge_page.
    # Preserve their rendered appearance instead of silently changing it.
    if source.get("/Annots") or source.get("/Group") or reader.trailer["/Root"].get("/OCProperties"):
        return False
    if source.rotation:
        retained = PdfWriter()
        source = retained.add_page(source)
        source.transfer_rotation_to_content()
    box = source.cropbox
    left, top, right, bottom = item["crop"]
    x0, y0, x1, y1 = map(float, box)
    box.lower_left = (x0 + left * (x1 - x0), y0 + bottom * (y1 - y0))
    box.upper_right = (x1 - right * (x1 - x0), y1 - top * (y1 - y0))
    x, y, w, h, scale = _placement(item, float(box.width), float(box.height))
    transform = Transformation().scale(scale).translate(x - float(box.left) * scale, page_height - y - h - float(box.bottom) * scale)
    target.merge_transformed_page(source, transform)
    return True


def _ppt_text(shapes, name, lines, x, y, w, h, size, family, color=(.08, .08, .09), leading=None):
    shape = shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    shape.name = name
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = family
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor(*(round(value * 255) for value in color))
        # Font.name sets only a:latin; CJK must also use the measured face.
        for tag in ("a:ea", "a:cs"):
            element = OxmlElement(tag)
            element.set("typeface", family)
            paragraph._p.get_or_add_pPr().get_or_add_defRPr().append(element)
        paragraph.line_spacing = Pt(leading if leading is not None else size * 1.2)
        paragraph.space_before = paragraph.space_after = Pt(0)
    return shape


@lru_cache(maxsize=32)
def _source_font(name):
    """Only known installed normal faces can safely become native source text."""
    from .drawings import _sheet_fonts
    font, family = _font()
    drawing = TTFont("AF_Publication_Source", str(_sheet_fonts()["normal"]))
    pdfmetrics.registerFont(drawing)
    clean = lambda value: re.sub(r"[^a-z0-9]", "", value.lower())
    name = re.sub(r"^[A-Z]{6}\+", "", name)
    aliases = {"Helvetica": "Arial", "ArialMT": "Arial"}
    name = aliases.get(name, name)
    for candidate in (font, drawing):
        family = candidate.face.familyName.decode("utf-8")
        names = (family, candidate.face.name.decode("utf-8"))
        if clean(name) in {clean(value) for value in names}:
            return candidate, family
    return None


def _ppt_pdf(slide, data, item):
    """Convert a bounded PDF subset, or refuse before adding any native shapes.

    Curves, clipping groups, transparency, dashed paths, image/shading content,
    rotated text and partly cropped primitives retain the exact raster preview.
    PDF export is independent of this bounded PPTX representation.
    """
    with fitz.open(stream=data, filetype="pdf") as document:
        page = document[item["source"]["pageIndex"]]
        if page.rotation or document.get_ocgs():
            return "rotated page or optional layers"
        if page.first_annot or page.first_widget:
            return "source annotations"
        if any(kind not in {"fill-path", "stroke-path", "fill-text"} for kind, _ in page.get_bboxlog()):
            return "images, shading or special text"
        rect = page.rect
        left, top, right, bottom = item["crop"]
        clip = fitz.Rect(rect.x0 + left * rect.width, rect.y0 + top * rect.height,
                         rect.x1 - right * rect.width, rect.y1 - bottom * rect.height)
        x, y, w, h, scale = _placement(item, clip.width, clip.height)
        entries = []
        for path in page.get_drawings(extended=True):
            if path["type"] not in {"s", "f", "fs"}:
                return "clipping or transparency group"
            if path.get("fill_opacity", 1) not in (None, 1) or path.get("stroke_opacity", 1) not in (None, 1):
                return "transparent paths"
            if path.get("dashes") not in (None, "[] 0"):
                return "dashed paths"
            if any(cap != 0 for cap in (path.get("lineCap") or (0,))) or path.get("lineJoin", 0) not in (None, 0):
                return "special line caps or joins"
            if any(command[0] not in {"l", "re"} for command in path["items"]):
                return "curved paths"
            bounds = fitz.Rect(path["rect"])
            pad = path.get("width") or 0
            bounds += (-pad / 2, -pad / 2, pad / 2, pad / 2)
            if not bounds.intersects(clip):
                continue
            if not clip.contains(bounds):
                return "partly cropped paths"
            commands = path["items"]
            connected = all(command[0] == "l" for command in commands) and all(commands[i-1][2] == commands[i][1] for i in range(1, len(commands)))
            if "f" in path["type"] and len(commands) > 1:
                if not connected or path.get("even_odd"):
                    return "compound filled paths"
            if path.get("closePath") and not connected and any(command[0] != "re" for command in commands):
                return "compound closed paths"
            if "f" in path["type"] and len(commands) == 1 and commands[0][0] == "l":
                return "degenerate filled path"
            entries.append((path["seqno"], "path", path))
        for span in page.get_texttrace():
            if span["type"] != 0 or span["opacity"] != 1 or span["dir"] != (1, 0) or span["wmode"] != 0:
                return "transformed or special text"
            if len(span["color"]) not in (1, 3):
                return "special text color"
            bounds = fitz.Rect(span["bbox"])
            if not bounds.intersects(clip):
                continue
            if not clip.contains(bounds):
                return "partly cropped text"
            if any(character[0] in (0, 0xfffd) for character in span["chars"]):
                return "unmapped text glyphs"
            match = _source_font(span["font"])
            if match is None:
                return "source font unavailable"
            font, family = match
            text = "".join(chr(char[0]) for char in span["chars"])
            advance = span["chars"][-1][3][2] - span["chars"][0][2][0]
            if abs(pdfmetrics.stringWidth(text, font.fontName, span["size"]) - advance) > max(.2, advance * .002):
                return "source text spacing"
            ascent, descent = pdfmetrics.getAscentDescent(font.fontName, span["size"])
            entries.append((span["seqno"], "text", span | {"family": family, "text": text, "font_ascent": ascent, "font_descent": descent}))
        if not entries:
            return "empty or unsupported source"

        def point(value):
            return x + (value[0] - clip.x0) * scale, y + (value[1] - clip.y0) * scale

        group = slide.shapes.add_group_shape()
        group.name = item["id"]
        for index, (_, kind, value) in enumerate(sorted(entries, key=lambda row: row[0])):
            name = f"{item['id']}-{index}"
            if kind == "text":
                origin = value["chars"][0][2]
                tx, baseline = point(origin)
                size = value["size"] * scale
                ascent, descent = value["font_ascent"] * scale, value["font_descent"] * scale
                color = value["color"]
                if len(color) == 1:
                    color = color * 3
                _ppt_text(group.shapes, name, [value["text"]], tx, baseline - ascent,
                          max(.1, (value["bbox"][2] - origin[0]) * scale), ascent - descent, size, value["family"], color, ascent - descent)
                continue
            commands = value["items"]
            created = []
            connected = all(command[0] == "l" for command in commands) and all(commands[i-1][2] == commands[i][1] for i in range(1, len(commands)))
            if len(commands) > 1 and connected:
                vertices = [point(commands[0][1]), *(point(command[2]) for command in commands)]
                builder = group.shapes.build_freeform(Pt(vertices[0][0]), Pt(vertices[0][1]))
                closed = value.get("closePath", False) or "f" in value["type"] or vertices[0] == vertices[-1]
                builder.add_line_segments([(Pt(px), Pt(py)) for px, py in vertices[1:]], close=closed)
                created.append(builder.convert_to_shape())
            else:
                for command in commands:
                    if command[0] == "re":
                        bounds = command[1]
                        sx, sy = point(bounds.top_left)
                        shape = group.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Pt(sx), Pt(sy), Pt(bounds.width * scale), Pt(bounds.height * scale))
                    else:
                        a, b = point(command[1]), point(command[2])
                        shape = group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(a[0]), Pt(a[1]), Pt(b[0]), Pt(b[1]))
                    created.append(shape)
            for shape in created:
                shape.name = name
                if "f" in value["type"]:
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = RGBColor(*(round(channel * 255) for channel in value["fill"]))
                elif hasattr(shape, "fill"):
                    shape.fill.background()
                if "s" in value["type"]:
                    shape.line.color.rgb = RGBColor(*(round(channel * 255) for channel in value["color"]))
                    shape.line.width = Pt((value["width"] or .1) * scale)
                else:
                    shape.line.fill.background()
        return None


def export_publication(binding, revision, format):
    publication = read_publication(binding, revision)
    if not publication["pages"]:
        raise StudioError(422, "PUBLICATION_EMPTY", "Add a page before exporting.")
    if any(row["status"] == "missing" for row in publication["sources"]):
        raise StudioError(409, "PUBLICATION_SOURCE_MISSING", "Restore the missing sources or remove their elements before exporting.")
    width, height = publication["spec"]["width"], publication["spec"]["height"]
    output = BytesIO()
    pdf = PdfWriter() if format == "pdf" else None
    ppt = Presentation() if format == "pptx" else None
    if pdf is not None:
        pdf.add_metadata({"/Title": publication["title"], "/Author": "MonkeyHub"})
    else:
        ppt.slide_width, ppt.slide_height = Pt(width), Pt(height)
        ppt.core_properties.title = publication["title"]
        ppt.core_properties.author = "MonkeyHub"
    for page in publication["pages"]:
        slide = ppt.slides.add_slide(ppt.slide_layouts[6]) if ppt is not None else None
        pdf_page = pdf.add_blank_page(width, height) if pdf is not None else None
        for item in page["elements"]:
            x, y, w, h = (item[key] for key in ("x", "y", "width", "height"))
            fragment = BytesIO()
            canvas = Canvas(fragment, pagesize=(width, height), invariant=1, pageCompression=1) if pdf is not None else None
            if item["kind"] == "image":
                source = item["source"]
                document, data = document_bytes(binding, source["runId"], source["assetSha256"], source.get("revisionRef"))
                fallback = None
                if document.mime_type == "application/pdf":
                    if pdf is not None:
                        if _pdf_source(pdf_page, data, item, height):
                            continue
                    else:
                        fallback = _ppt_pdf(slide, data, item)
                        if fallback is None:
                            continue
                data, (x, y, w, h) = _image(data, document.mime_type, item)
                if canvas is not None:
                    canvas.drawImage(ImageReader(BytesIO(data)), x, height - y - h, width=w, height=h, mask="auto")
                else:
                    picture = slide.shapes.add_picture(BytesIO(data), Pt(x), Pt(y), Pt(w), Pt(h))
                    picture.name = item["id"] if fallback is None else f"{item['id']} (raster preview: {fallback})"
                    if fallback:
                        picture._element.nvPicPr.cNvPr.set("descr", f"Retained PDF raster preview: {fallback}. PDF export preserves compatible source vectors.")
            else:
                font, family = _font()
                size = item["fontSize"]
                lines = _lines(item["text"], w, size)
                ascent, descent = pdfmetrics.getAscentDescent(font.fontName, size)
                leading = max(size * 1.2, ascent - descent)
                if ascent - descent + (len(lines) - 1) * leading > h:
                    raise StudioError(422, "PUBLICATION_TEXT_OVERFLOW", f"Text in {item['id']} does not fit. Enlarge the box or shorten the text.")
                if canvas is not None:
                    canvas.setFillColorRGB(.08, .08, .09)
                    canvas.setFont(font.fontName, size)
                    for index, line in enumerate(lines):
                        canvas.drawString(x, height - y - ascent - index * leading, line)
                else:
                    _ppt_text(slide.shapes, item["id"], lines, x, y, w, h, size, family, leading=leading)
            if canvas is not None:
                canvas.save()
                pdf_page.merge_page(PdfReader(BytesIO(fragment.getvalue())).pages[0])
    if pdf is not None:
        pdf.write(output)
    else:
        ppt.save(output)
    name = re.sub(r"[^\w\-\u4e00-\u9fff]+", "-", publication["title"]).strip("-") or "publication"
    return BoardExport(name + "." + format,
        "application/pdf" if pdf is not None else "application/vnd.openxmlformats-officedocument.presentationml.presentation", output.getvalue())
