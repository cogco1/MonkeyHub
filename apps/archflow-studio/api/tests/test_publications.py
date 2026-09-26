"""Editable publication behavior through HTTP and real retained project files."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from io import BytesIO
import base64
import unittest

import fitz
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas

from . import test_boards as board_helpers
from .test_documents import image_bytes, two_page_pdf
from .support import PROJECT_ID


def source(document, page=0):
    return {"runId": document["runId"], "assetSha256": document["assetSha256"],
            "revisionRef": document.get("revisionRef"), "pageIndex": page}


def page(identifier, title="Spatial sequence", document=None):
    elements = [{"id": identifier + "-title", "kind": "text", "x": 20, "y": 20,
                 "width": 520, "height": 60, "text": title, "fontSize": 20}]
    if document:
        elements.append({"id": identifier + "-image", "kind": "image", "x": 20, "y": 100,
                         "width": 520, "height": 260, "source": source(document)})
    return {"id": identifier, "elements": elements}


def request(pages, base=None, title="Design review"):
    return {"projectId": PROJECT_ID, "baseRevisionSha256": base, "title": title,
            "spec": {"width": 600, "height": 400, "template": "hero"}, "pages": pages}


class PublicationTests(unittest.TestCase):
    # Share the small real-project fixture without inheriting the Board suite.
    setUp = board_helpers.BoardTests.setUp
    new_client = board_helpers.BoardTests.new_client
    files = board_helpers.BoardTests.files
    upload = board_helpers.BoardTests.upload

    def save_publication(self, body):
        response = self.client.put("/api/publication", json=body)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def export(self, revision, format="pdf"):
        return self.client.post("/api/publication/export", json={
            "projectId": PROJECT_ID, "revisionSha256": revision, "format": format})

    def assert_design_unchanged(self):
        self.assertEqual(self.repository.read_head(), self.head)
        self.assertEqual(self.repository.read_design_branches(), {})

    def test_empty_read_and_invalid_first_save_do_not_write(self):
        before = self.files()
        response = self.client.get("/api/publication")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["pages"], [])
        self.assertIsNone(response.json()["revisionSha256"])
        bad = page("outside")
        bad["elements"][0]["x"] = 590
        self.assertEqual(self.client.put("/api/publication", json=request([bad])).status_code, 422)
        self.assertEqual(self.files(), before)
        self.assert_design_unchanged()

    def test_order_text_and_exact_image_survive_reopen_and_idempotent_save(self):
        image = self.upload(image_bytes(), "section.png", "image/png")
        saved = self.save_publication(request([page("second", "Second first", image), page("first", "First last")]))
        before = self.files()
        reopened = self.new_client().get("/api/publication").json()
        self.assertEqual(reopened, saved)
        self.assertEqual([row["id"] for row in reopened["pages"]], ["second", "first"])
        self.assertEqual(reopened["pages"][0]["elements"][1]["source"], source(image))
        unchanged = self.save_publication(request(saved["pages"], saved["revisionSha256"]))
        self.assertEqual(unchanged, saved)
        self.assertEqual(self.files(), before)
        self.assert_design_unchanged()

    def test_concurrent_saves_have_one_winner_and_history_stays_readable(self):
        first = self.save_publication(request([page("p1")]))
        client_a, client_b = self.new_client(), self.new_client()
        bodies = [request([page("p1", title)], first["revisionSha256"]) for title in ("Option A", "Option B")]
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda pair: pair[0].put("/api/publication", json=pair[1]), zip((client_a, client_b), bodies)))
        self.assertEqual(sorted(row.status_code for row in results), [200, 409])
        winner = next(row.json() for row in results if row.status_code == 200)
        self.assertEqual(self.new_client().get("/api/publication").json(), winner)
        self.assertEqual(self.client.get("/api/publication", params={"revision": first["revisionSha256"]}).json(), first)
        self.assert_design_unchanged()

    def test_cross_project_requests_and_unknown_revisions_are_rejected_without_writes(self):
        saved = self.save_publication(request([page("p1")]))
        before = self.files()
        bad = request([page("p1", "Overwrite")], saved["revisionSha256"]) | {"projectId": "other-project"}
        self.assertEqual(self.client.put("/api/publication", json=bad).status_code, 403)
        self.assertEqual(self.client.post("/api/publication/export", json={"projectId": "other-project",
            "revisionSha256": saved["revisionSha256"], "format": "pdf"}).status_code, 403)
        self.assertEqual(self.client.post("/api/publication/from-board", json={"projectId": "other-project",
            "baseRevisionSha256": saved["revisionSha256"], "boardRevisionSha256": "a" * 64, "elementIds": ["any"]}).status_code, 403)
        self.assertEqual(self.client.get("/api/publication", params={"revision": "b" * 64}).status_code, 404)
        self.assertEqual(self.export("b" * 64).status_code, 404)
        self.assertEqual(self.files(), before)

    def test_pdf_is_repeatable_and_pptx_has_editable_text_with_matching_page_order(self):
        image = self.upload(image_bytes(color="blue"), "section.png", "image/png")
        saved = self.save_publication(request([page("p2", "Review title\nEditable judgment", image), page("p1", "Second page")]))
        before = self.files()
        pdf = self.export(saved["revisionSha256"])
        self.assertEqual(pdf.status_code, 200, pdf.text[:200] if pdf.status_code != 200 else "")
        self.assertEqual(pdf.content, self.export(saved["revisionSha256"]).content)
        reader = PdfReader(BytesIO(pdf.content))
        self.assertEqual(len(reader.pages), 2)
        self.assertIn("Editable judgment", reader.pages[0].extract_text())
        self.assertIn("Second page", reader.pages[1].extract_text())
        self.assertEqual(tuple(map(float, reader.pages[0].mediabox)), (0, 0, 600, 400))
        ppt = self.export(saved["revisionSha256"], "pptx")
        self.assertEqual(ppt.status_code, 200)
        deck = Presentation(BytesIO(ppt.content))
        self.assertEqual(len(deck.slides), 2)
        title = next(shape for shape in deck.slides[0].shapes if shape.has_text_frame)
        self.assertEqual(title.text, "Review title\nEditable judgment")
        title.text_frame.paragraphs[1].text = "Changed in PowerPoint"
        edited = BytesIO(); deck.save(edited)
        self.assertIn("Changed in PowerPoint", Presentation(BytesIO(edited.getvalue())).slides[0].shapes[0].text)
        picture = next(shape for shape in deck.slides[0].shapes if hasattr(shape, "image"))
        with Image.open(BytesIO(picture.image.blob)) as pixels:
            self.assertEqual(pixels.getpixel((pixels.width // 2, pixels.height // 2))[:3], (0, 0, 255))
        self.assertEqual(self.files(), before, "exports do not mint project content or advance state")
        self.assert_design_unchanged()

    def test_replacement_marks_old_source_stale_and_freezing_preserves_original_export(self):
        original = self.upload(image_bytes(color="blue"), "section.png", "image/png")
        saved = self.save_publication(request([page("p1", document=original)]))
        baseline = self.export(saved["revisionSha256"]).content
        replacement = self.client.post("/api/documents", json={"projectId": PROJECT_ID, "fileName": "section-revised.png",
            "mimeType": "image/png", "contentBase64": base64.b64encode(image_bytes(color="red")).decode(),
            "replacesPages": [source(original) | {"newPageIndex": 0}]})
        self.assertEqual(replacement.status_code, 201, replacement.text)
        revised = replacement.json()
        stale = self.new_client().get("/api/publication").json()
        self.assertEqual(stale["sources"][0]["status"], "stale")
        self.assertEqual(stale["sources"][0]["replacement"], source(revised))
        self.assertEqual(stale["pages"], saved["pages"], "source changes must not silently rewrite composed pages")
        frozen_pages = deepcopy(saved["pages"])
        frozen_pages[0]["elements"][1]["frozen"] = True
        frozen = self.save_publication(request(frozen_pages, saved["revisionSha256"]))
        self.assertEqual(frozen["sources"][0]["status"], "frozen")
        self.assertEqual(self.export(frozen["revisionSha256"]).content, baseline)
        updated_pages = deepcopy(frozen_pages)
        updated_pages[0]["elements"][1].update(source=source(revised), frozen=False)
        updated = self.save_publication(request(updated_pages, frozen["revisionSha256"]))
        self.assertEqual(updated["sources"][0]["status"], "current")
        self.assertNotEqual(self.export(updated["revisionSha256"]).content, baseline)
        self.assert_design_unchanged()

    def test_missing_retained_source_allows_text_repair_but_blocks_export_until_restored(self):
        original = self.upload(image_bytes(), "section.png", "image/png")
        saved = self.save_publication(request([page("p1", document=original)]))
        digest = original["assetSha256"]
        blob = self.repository.layout.resolve_relative(f"objects/sha256/{digest[:2]}/{digest}")
        original_bytes = blob.read_bytes()
        blob.write_bytes(b"damaged retained file")
        missing = self.new_client().get("/api/publication")
        self.assertEqual(missing.status_code, 200, missing.text)
        self.assertEqual(missing.json()["sources"][0]["status"], "missing")
        edited = deepcopy(saved["pages"])
        edited[0]["elements"][0]["text"] = "Waiting for source recovery"
        repaired = self.save_publication(request(edited, saved["revisionSha256"]))
        self.assertEqual(self.export(repaired["revisionSha256"]).status_code, 409)
        blob.write_bytes(original_bytes)
        self.assertEqual(self.new_client().get("/api/publication").json()["sources"][0]["status"], "current")
        self.assertEqual(self.export(repaired["revisionSha256"]).status_code, 200)
        self.assert_design_unchanged()

    def test_unknown_image_or_pdf_page_cannot_be_introduced(self):
        document = self.upload(two_page_pdf())
        invalid = page("p1", document=document)
        invalid["elements"][1]["source"]["pageIndex"] = 2
        before = self.files()
        self.assertEqual(self.client.put("/api/publication", json=request([invalid])).status_code, 422)
        invalid["elements"][1]["source"].update(pageIndex=0, assetSha256="0" * 64)
        self.assertEqual(self.client.put("/api/publication", json=request([invalid])).status_code, 422)
        invalid["elements"][1]["source"] = source(document) | {"revisionRef": "unretained-revision"}
        self.assertEqual(self.client.put("/api/publication", json=request([invalid])).status_code, 422)
        self.assertEqual(self.files(), before)

    def test_board_selection_uses_named_saved_revision_and_frame_order_only(self):
        document = self.upload(two_page_pdf())
        later = board_helpers.image_element(document, 1, "later") | {"x": 30, "y": 300, "frameId": "review-frame"}
        earlier = board_helpers.image_element(document, 0, "earlier") | {"x": 30, "y": 30, "frameId": "review-frame"}
        unselected = board_helpers.image_element(document, 0, "unselected")
        removed = board_helpers.image_element(document, 0, "deleted") | {"isDeleted": True, "frameId": "review-frame"}
        frame = {"id": "review-frame", "type": "frame", "x": 0, "y": 0, "width": 800, "height": 700, "name": "Selected pages"}
        board = self.client.put("/api/board", json=board_helpers.body([later, unselected, frame, earlier, removed]))
        self.assertEqual(board.status_code, 200, board.text)
        revision = board.json()["revisionSha256"]
        changed = self.client.put("/api/board", json=board_helpers.body([unselected], revision))
        self.assertEqual(changed.status_code, 200, changed.text)
        appended = self.client.post("/api/publication/from-board", json={"projectId": PROJECT_ID,
            "baseRevisionSha256": None, "boardRevisionSha256": revision, "elementIds": ["review-frame"]})
        self.assertEqual(appended.status_code, 200, appended.text)
        saved = appended.json()
        self.assertEqual(len(saved["pages"]), 2)
        sources = [element["source"] for row in saved["pages"] for element in row["elements"] if element["kind"] == "image"]
        self.assertEqual(sources, [source(document, 0), source(document, 1)])
        self.assertEqual(self.client.get("/api/board").json(), changed.json())
        self.assertEqual(self.new_client().get("/api/publication").json(), saved)
        before = self.files()
        invalid = self.client.post("/api/publication/from-board", json={"projectId": PROJECT_ID,
            "baseRevisionSha256": saved["revisionSha256"], "boardRevisionSha256": revision, "elementIds": ["not-saved"]})
        self.assertEqual(invalid.status_code, 409, invalid.text)
        self.assertEqual(self.files(), before)
        exported = self.export(saved["revisionSha256"])
        self.assertEqual(exported.status_code, 200, exported.text)
        self.assertEqual(len(PdfReader(BytesIO(exported.content)).pages), 2)
        self.assert_design_unchanged()

    def test_text_overflow_is_reported_for_both_formats_without_writes(self):
        content = page("p1", "This text cannot fit in a single short line")
        content["elements"][0].update(width=45, height=25, fontSize=20)
        saved = self.save_publication(request([content]))
        before = self.files()
        for format in ("pdf", "pptx"):
            result = self.export(saved["revisionSha256"], format)
            self.assertEqual(result.status_code, 422, result.text)
            self.assertIn("PUBLICATION_TEXT_OVERFLOW", result.text)
        self.assertEqual(self.files(), before)

    def test_explicit_board_order_reuses_the_rule_and_deduplicates_frame_children(self):
        document = self.upload(two_page_pdf())
        top = board_helpers.image_element(document, 0, "top") | {"y": 20, "frameId": "frame"}
        bottom = board_helpers.image_element(document, 1, "bottom") | {"y": 200, "frameId": "frame"}
        frame = {"id": "frame", "type": "frame", "x": 0, "y": 0, "width": 500, "height": 500}
        board = self.client.put("/api/board", json=board_helpers.body([frame, top, bottom])).json()
        body = {"projectId": PROJECT_ID, "baseRevisionSha256": None,
                "boardRevisionSha256": board["revisionSha256"], "elementIds": ["bottom", "frame", "top"]}
        result = self.client.post("/api/publication/from-board", json=body)
        self.assertEqual(result.status_code, 200, result.text)
        saved = result.json()
        self.assertEqual([row["elements"][1]["source"] for row in saved["pages"]], [source(document, 1), source(document, 0)])
        geometry = lambda row: [{key: element[key] for key in ("kind", "x", "y", "width", "height", "fontSize")} for element in row["elements"]]
        self.assertEqual(geometry(saved["pages"][0]), geometry(saved["pages"][1]))
        body["baseRevisionSha256"] = saved["revisionSha256"]
        self.assertEqual(self.client.post("/api/publication/from-board", json=body).json(), saved)
        self.assertEqual(self.new_client().get("/api/publication").json(), saved)
        self.assert_design_unchanged()

    def test_image_crop_is_identical_in_pdf_and_pptx_and_preserves_source_bytes(self):
        pixels = Image.new("RGB", (200, 100), "red")
        pixels.paste("blue", (100, 0, 200, 100))
        original = BytesIO(); pixels.save(original, format="PNG")
        document = self.upload(original.getvalue(), "two-halves.png", "image/png")
        content = page("p1", document=document)
        content["elements"][1].update(crop=[.5, 0, 0, 0], x=20, y=100, width=400, height=200)
        saved = self.save_publication(request([content]))
        before = self.files()
        pdf = PdfReader(BytesIO(self.export(saved["revisionSha256"]).content))
        deck = Presentation(BytesIO(self.export(saved["revisionSha256"], "pptx").content))
        picture = next(shape for shape in deck.slides[0].shapes if hasattr(shape, "image"))
        exported_images = [pdf.pages[0].images[0].image, Image.open(BytesIO(picture.image.blob))]
        for image in exported_images:
            self.assertEqual(image.size, (100, 100))
            self.assertEqual(image.convert("RGB").getextrema(), ((0, 0), (0, 0), (255, 255)))
        self.assertEqual((picture.left.pt, picture.top.pt, picture.width.pt, picture.height.pt), (120, 100, 200, 200))
        self.assertEqual(self.client.get(f"/api/documents/{document['assetSha256']}/bytes", params={"runId": document["runId"]}).content, original.getvalue())
        self.assertEqual(self.files(), before)

    def test_source_drawing_keeps_native_vectors_and_editable_pptx_objects(self):
        from archflow_studio_api.application.drawings import _sheet_fonts
        from monkeydiagram.drawing_output import PaperCanvas, render_pdf, MM_PER_PT
        drawing = PaperCanvas(font_mapping={"normal": _sheet_fonts()["normal"]})
        drawing.start_sheet("A01", (200 * MM_PER_PT, 100 * MM_PER_PT))
        drawing.setFont("normal", 12)
        drawing.rect(10, 10, 150, 70)
        drawing.line(10.25, 12.5, 50.75, 12.5)
        drawing.drawString(20, 40, "Editable drawing")
        path = drawing.beginPath()
        path.moveTo(10.25, 15.5)
        path.lineTo(50.75, 15.5)
        path.lineTo(30.25, 40.125)
        path.close()
        drawing.drawPath(path)
        drawing.showPage()
        original = render_pdf(drawing)
        document = self.upload(original, "drawing.pdf", "application/pdf")
        saved = self.save_publication(request([page("native", document=document)]))
        before = self.files()
        pdf = self.export(saved["revisionSha256"])
        self.assertEqual(pdf.status_code, 200, pdf.text if pdf.status_code != 200 else "")
        self.assertEqual(pdf.content, self.export(saved["revisionSha256"]).content)
        with fitz.open(stream=pdf.content, filetype="pdf") as parsed:
            self.assertIn("Editable drawing", parsed[0].get_text())
            self.assertGreaterEqual(len(parsed[0].get_drawings()), 3)
            self.assertEqual(parsed[0].get_images(), [])
            rect = next(row["rect"] for row in parsed[0].get_drawings() if row["items"][0][0] == "re")
            self.assertAlmostEqual(rect.x0, 46, places=3)
            self.assertAlmostEqual(rect.y0, 152, places=3)
            self.assertAlmostEqual(rect.width, 390, places=3)
        ppt = self.export(saved["revisionSha256"], "pptx")
        self.assertEqual(ppt.status_code, 200, ppt.text if ppt.status_code != 200 else "")
        deck = Presentation(BytesIO(ppt.content))
        group = next(shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP)
        native_text = next(shape for shape in group.shapes if shape.has_text_frame and shape.text == "Editable drawing")
        rectangle = next(shape for shape in group.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
        self.assertAlmostEqual(rectangle.width.pt, 390, places=3)
        triangle = next(shape for shape in group.shapes if shape.shape_type == MSO_SHAPE_TYPE.FREEFORM)
        self.assertTrue(triangle._element.xpath(".//a:close"))
        self.assertAlmostEqual(triangle.left.pt, 46.65, places=3, msg="Fractional paper coordinates must not round to whole points")
        native_text.text = "Changed drawing label"
        rectangle.width = Pt(321)
        rewritten = BytesIO()
        deck.save(rewritten)
        reopened = Presentation(BytesIO(rewritten.getvalue()))
        edited_group = next(shape for shape in reopened.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP)
        self.assertTrue(any(shape.has_text_frame and shape.text == "Changed drawing label" for shape in edited_group.shapes))
        self.assertEqual(next(shape for shape in edited_group.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE).width.pt, 321)
        self.assertEqual(self.files(), before)
        self.assert_design_unchanged()

    def test_pdf_crop_transparency_and_element_order_survive_vector_composition(self):
        source_pdf = BytesIO()
        canvas = Canvas(source_pdf, pagesize=(200, 100), invariant=1)
        canvas.setFillColorRGB(0, 0, 1)
        canvas.setFillAlpha(.5)
        canvas.rect(100, 0, 100, 100, stroke=0, fill=1)
        canvas.save()
        document = self.upload(source_pdf.getvalue(), "transparent.pdf", "application/pdf")
        red = self.upload(image_bytes(color="red"), "background.png", "image/png")
        content = page("overlay", document=red)
        content["elements"][1].update(x=20, y=100, width=200, height=100)
        content["elements"].append({"id": "vector", "kind": "image", "x": 20, "y": 100,
            "width": 200, "height": 100, "source": source(document)})
        content["elements"].append({"id": "top", "kind": "text", "x": 130, "y": 130,
            "width": 80, "height": 30, "text": "TOP", "fontSize": 20})
        saved = self.save_publication(request([content]))
        pdf = self.export(saved["revisionSha256"])
        self.assertEqual(pdf.content, self.export(saved["revisionSha256"]).content)
        with fitz.open(stream=pdf.content, filetype="pdf") as parsed:
            raster = parsed[0].get_pixmap(alpha=False)
            pixels = Image.frombytes("RGB", (raster.width, raster.height), raster.samples)
            # The square PNG is contained in the wider box, at x=70..170.
            self.assertEqual(pixels.getpixel((90, 180)), (255, 0, 0))
            r, g, b = pixels.getpixel((150, 180))
            self.assertTrue(120 <= r <= 135 and g == 0 and 120 <= b <= 135)
            self.assertEqual(parsed[0].get_bboxlog()[-1][0], "fill-text")
            self.assertEqual(len(parsed[0].get_drawings()), 1)
        ppt = Presentation(BytesIO(self.export(saved["revisionSha256"], "pptx").content))
        preview = next(shape for shape in ppt.slides[0].shapes if "raster preview" in shape.name)
        self.assertIn("transparent paths", preview.name)
        with Image.open(BytesIO(preview.image.blob)) as image:
            self.assertEqual(image.mode, "RGBA")
            self.assertEqual(image.getpixel((10, 10))[3], 0)
        content["elements"][2].update(crop=[.5, 0, 0, 0], x=300, y=100, width=200, height=200)
        cropped = self.save_publication(request([content], saved["revisionSha256"]))
        with fitz.open(stream=self.export(cropped["revisionSha256"]).content, filetype="pdf") as parsed:
            row = parsed[0].get_drawings()[0]
            self.assertAlmostEqual(row["rect"].x0, 300)
            self.assertAlmostEqual(row["rect"].width, 200)
            self.assertAlmostEqual(row["rect"].height, 200)
        self.assert_design_unchanged()

    def test_unsupported_source_font_and_curve_are_explicit_raster_fallbacks(self):
        for name, drawing, reason in (
            ("font", lambda c: (c.setFont("Courier", 12), c.drawString(20, 40, "Native font unavailable")), "source font unavailable"),
            ("curve", lambda c: c.circle(50, 50, 20), "curved paths"),
        ):
            with self.subTest(name=name):
                data = BytesIO()
                canvas = Canvas(data, pagesize=(200, 100), invariant=1)
                drawing(canvas)
                canvas.save()
                document = self.upload(data.getvalue(), name + ".pdf", "application/pdf")
                current = self.client.get("/api/publication").json()
                saved = self.save_publication(request([page(name, document=document)], current["revisionSha256"]))
                result = self.export(saved["revisionSha256"], "pptx")
                self.assertEqual(result.status_code, 200, result.text if result.status_code != 200 else "")
                deck = Presentation(BytesIO(result.content))
                picture = next(shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
                self.assertIn(reason, picture.name)
                self.assertIn(reason, picture._element.nvPicPr.cNvPr.get("descr"))
                with fitz.open(stream=self.export(saved["revisionSha256"]).content, filetype="pdf") as parsed:
                    self.assertEqual(parsed[0].get_images(), [])

    def test_embedded_font_shared_line_widths_and_page_bottom_do_not_clip(self):
        from archflow_studio_api.application.publication_output import _font
        from fontTools.ttLib import TTFont
        font, family = _font()
        content = page("type")
        content["elements"][0].update(text="WWWWWWW", width=140, height=60, fontSize=20)
        content["elements"].append({"id": "bottom", "kind": "text", "x": 10, "y": 376,
            "width": 200, "height": 24, "text": "gypqj", "fontSize": 20})
        self.assertTrue(all(ord(char) in font.face.charToGlyph for char in "中文"))
        content["elements"].append({"id": "cjk", "kind": "text", "x": 20, "y": 100,
            "width": 200, "height": 30, "text": "中文", "fontSize": 20})
        saved = self.save_publication(request([content]))
        pdf = self.export(saved["revisionSha256"])
        self.assertEqual(pdf.status_code, 200)
        reader = PdfReader(BytesIO(pdf.content))
        actual_fonts = [row.get_object() for row in reader.pages[0]["/Resources"]["/Font"].values()]
        self.assertTrue(any("/FontFile2" in row.get("/FontDescriptor", {}) for row in actual_fonts))
        with fitz.open(stream=pdf.content, filetype="pdf") as parsed:
            spans = [span for block in parsed[0].get_text("dict")["blocks"] for line in block.get("lines", []) for span in line["spans"]]
            self.assertTrue(any(span["text"] == "gypqj" for span in spans))
            self.assertLessEqual(max(span["bbox"][3] for span in spans), 400)
            self.assertIn("中文", parsed[0].get_text())
        deck = Presentation(BytesIO(self.export(saved["revisionSha256"], "pptx").content))
        title = deck.slides[0].shapes[0]
        # Independently read installed glyph advances, rather than asserting the
        # compiler's own width function or hard-coding one operating system font.
        with TTFont(font.face.filename, fontNumber=0) as measured:
            cmap = measured.getBestCmap()
            widths = measured["hmtx"].metrics
            units = measured["head"].unitsPerEm
            for paragraph in title.text_frame.paragraphs:
                self.assertEqual(paragraph.font.name, family)
                advance = sum(widths[cmap[ord(char)]][0] for char in paragraph.text) * 20 / units
                self.assertLessEqual(advance, title.width.pt)
        for shape in deck.slides[0].shapes:
            self.assertEqual(shape._element.xpath(".//a:ea")[0].get("typeface"), family)

    def test_single_oversize_or_unmapped_glyph_refuses_both_exports(self):
        for text, width, code in (("W", 1, "PUBLICATION_TEXT_OVERFLOW"), ("\U0010ffff", 200, "PUBLICATION_FONT_GLYPH")):
            with self.subTest(code=code):
                content = page("glyph", text)
                content["elements"][0]["width"] = width
                current = self.client.get("/api/publication").json()
                saved = self.save_publication(request([content], current["revisionSha256"]))
                for format in ("pdf", "pptx"):
                    exported = self.export(saved["revisionSha256"], format)
                    self.assertEqual(exported.status_code, 422, exported.text)
                    self.assertIn(code, exported.text)

    def test_annotated_pdf_uses_visual_fallback_and_crops_annotations_with_content(self):
        with fitz.open() as original:
            source_page = original.new_page(width=200, height=100)
            source_page.insert_text((10, 20), "Annotated source")
            annotation = source_page.add_rect_annot(fitz.Rect(120, 20, 180, 80))
            annotation.set_colors(stroke=(1, 0, 0), fill=(1, 0, 0))
            annotation.set_border(width=0)
            annotation.update(opacity=1)
            data = original.tobytes(no_new_id=True)
        document = self.upload(data, "annotated.pdf", "application/pdf")
        content = page("annotation", document=document)
        content["elements"][1].update(width=400, height=200)
        saved = self.save_publication(request([content]))
        pdf = self.export(saved["revisionSha256"])
        self.assertEqual(pdf.content, self.export(saved["revisionSha256"]).content)
        with fitz.open(stream=pdf.content, filetype="pdf") as parsed:
            self.assertEqual(len(parsed[0].get_images()), 1)
            self.assertIsNone(parsed[0].first_annot)
            pixmap = parsed[0].get_pixmap(alpha=False)
            pixels = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            self.assertEqual(pixels.getpixel((320, 200)), (255, 0, 0))
        deck = Presentation(BytesIO(self.export(saved["revisionSha256"], "pptx").content))
        self.assertTrue(any("raster preview: source annotations" in shape.name for shape in deck.slides[0].shapes))
        # Crop to the left half; an annotation lives outside the PDF content
        # stream and would otherwise survive merge_page's content-only clip.
        content["elements"][1]["crop"] = [0, 0, .5, 0]
        cropped = self.save_publication(request([content], saved["revisionSha256"]))
        with fitz.open(stream=self.export(cropped["revisionSha256"]).content, filetype="pdf") as parsed:
            pixmap = parsed[0].get_pixmap(alpha=False)
            pixels = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            self.assertFalse(any(red > 200 and green < 50 and blue < 50 for _, (red, green, blue) in pixels.getcolors(pixels.width * pixels.height)))
        self.assert_design_unchanged()
